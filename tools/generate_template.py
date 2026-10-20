"""
Fleet ERP — Database Template Generator
=======================================
Builds FLEET_ERP_DATABASE_TEMPLATE.xlsx: the single "database" workbook the user
uploads to the app. Every sheet carries its real column headers, dropdown
validation on constrained fields, worked example rows for two aircraft types
(A320neo + ATR72), and a mixed-method financial scenario demonstrating the
Scenario -> Rules -> Method + Params routing.

Run:  python3 tools/generate_template.py
Output: data/FLEET_ERP_DATABASE_TEMPLATE.xlsx
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ─────────────────────────────────────────────────────────────
# Controlled vocabularies (dropdown sources)
# ─────────────────────────────────────────────────────────────
ENUMS: dict[str, list[str]] = {
    "Category":                 ["AIRFRAME", "ENGINE", "APU", "LLP", "LANDING_GEAR"],
    "Driver":                   ["FH", "FC", "CALENDAR_MONTHS", "CALENDAR_DAYS"],
    "SoftThresholdType":        ["PCT", "ABS_FH", "ABS_FC", "ABS_MONTHS", "NONE"],
    "RedeliveryMinLifeType":    ["FH", "FC", "MONTHS", "NONE"],
    "Bool":                     ["TRUE", "FALSE"],
    "WorkscopeClass":           ["LIGHT", "HEAVY", "LLP_REPLACEMENT", "NONE"],
    "LinkType":                 ["HARD", "SOFT_BUNDLE"],
    "AircraftStatus":           ["ACTIVE", "AOG", "STORED"],
    "ProjectionEndMode":        ["REDELIVERY", "FIXED_DATE", "HORIZON_MONTHS"],
    "ScopeType":                ["ALL", "CATEGORY", "COMPONENT"],
    "MethodType":               ["MR_STANDARD", "MR_FH_FC_RATIO", "PBH_FLAT",
                                 "TM_WORKSCOPE", "EOL_PRORATA"],
    "RateBasis":                ["PER_FH", "PER_FC", "PER_MONTH"],
    "CompBasis":                ["PER_REMAINING_FH", "PER_CONSUMED_FH",
                                 "PER_REMAINING_FC", "PER_CONSUMED_FC"],
    "EventType":                ["LIGHT_SV", "HEAVY_SV", "LLP_REPLACEMENT"],
}

# ─────────────────────────────────────────────────────────────
# Styling
# ─────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
EXAMPLE_FONT = Font(italic=True, color="7F7F7F")
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
NOTE_FONT = Font(size=10, color="404040")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def _style_header(ws, ncols: int):
    for ci in range(1, ncols + 1):
        c = ws.cell(row=1, column=ci)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def _autoscale(ws, headers: list[str], min_w: int = 12, max_w: int = 34):
    for ci, h in enumerate(headers, start=1):
        width = max(min_w, min(max_w, len(str(h)) + 4))
        ws.column_dimensions[get_column_letter(ci)].width = width


def _add_dropdown(ws, col_idx: int, enum_key: str, first_row: int = 2, last_row: int = 400):
    """Attach a list validation pointing at the hidden _Lists sheet."""
    values = ENUMS[enum_key]
    # Lists live in _Lists; find the column reserved for this enum.
    col_letter = get_column_letter(_LIST_COL[enum_key])
    rng = f"_Lists!${col_letter}$2:${col_letter}${1 + len(values)}"
    dv = DataValidation(type="list", formula1=rng, allow_blank=True)
    dv.error = f"Pick one of: {', '.join(values)}"
    dv.errorTitle = "Invalid value"
    dv.prompt = f"Allowed: {', '.join(values)}"
    dv.promptTitle = enum_key
    ws.add_data_validation(dv)
    cl = get_column_letter(col_idx)
    dv.add(f"{cl}{first_row}:{cl}{last_row}")


# Reserve one column per enum on the hidden _Lists sheet
_LIST_COL: dict[str, int] = {k: i + 1 for i, k in enumerate(ENUMS.keys())}


def _write_rows(ws, headers: list[str], rows: list[list], example_rows: int = 0):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    _style_header(ws, len(headers))
    _autoscale(ws, headers)
    # Mark example rows in grey italic so the user knows what to delete
    for ri in range(2, 2 + example_rows):
        for ci in range(1, len(headers) + 1):
            cell = ws.cell(row=ri, column=ci)
            cell.font = EXAMPLE_FONT
            cell.border = BORDER


# ─────────────────────────────────────────────────────────────
# Build
# ─────────────────────────────────────────────────────────────

def build() -> Path:
    wb = Workbook()

    # ---- Hidden _Lists sheet (dropdown sources) ----
    lists = wb.active
    lists.title = "_Lists"
    for key, col in _LIST_COL.items():
        lists.cell(row=1, column=col, value=key).font = Font(bold=True)
        for i, val in enumerate(ENUMS[key], start=2):
            lists.cell(row=i, column=col, value=val)
    lists.sheet_state = "hidden"

    # ---- Instructions ----
    ins = wb.create_sheet("Instructions")
    _build_instructions(ins)

    # ---- Config ----
    cfg = wb.create_sheet("Config")
    _write_rows(
        cfg,
        ["Setting", "Value", "Notes"],
        [
            ["As_Of_Date", "2026-01-31", "Snapshot date the projection starts from (YYYY-MM-DD, month-end)."],
            ["Projection_End_Mode", "REDELIVERY", "REDELIVERY | FIXED_DATE | HORIZON_MONTHS"],
            ["Projection_End_Value", "", "Date (if FIXED_DATE) or month count (if HORIZON_MONTHS). Blank for REDELIVERY."],
            ["Base_Currency", "USD", "Reporting currency for all financial output."],
            ["Schema_Version", "1.0", "Do not edit — used by the app to validate the file."],
        ],
    )
    _add_dropdown(cfg, 2, "ProjectionEndMode", first_row=3, last_row=3)
    cfg.column_dimensions["C"].width = 70

    # ---- REG_<family> : component / event registry (one per type) ----
    reg_headers = [
        "Component_ID", "Component_Name", "Category", "Driver",
        "Hard_Limit", "Soft_Threshold_Type", "Soft_Threshold_Value",
        "Redelivery_Min_Life_Type", "Redelivery_Min_Life_Value",
        "Track_TSN_CSN", "Workscope_Class", "SV_Event_Label",
    ]
    reg_a320 = [
        ["AF_6Y", "Airframe 6-Year Check", "AIRFRAME", "CALENDAR_MONTHS", 72, "PCT", 0.85, "MONTHS", 6, "FALSE", "HEAVY", "6Y Check"],
        ["AF_12Y", "Airframe 12-Year Check", "AIRFRAME", "CALENDAR_MONTHS", 144, "PCT", 0.85, "MONTHS", 6, "FALSE", "HEAVY", "12Y Check"],
        ["ENG1_PR", "Engine 1 Performance Restoration", "ENGINE", "FH", 20000, "PCT", 0.80, "FH", 3000, "TRUE", "HEAVY", "ENG1 PR SV"],
        ["ENG2_PR", "Engine 2 Performance Restoration", "ENGINE", "FH", 20000, "PCT", 0.80, "FH", 3000, "TRUE", "HEAVY", "ENG2 PR SV"],
        ["ENG1_LLP_HPT", "Engine 1 HPT Disk (LLP)", "LLP", "FC", 18000, "PCT", 0.90, "FC", 1500, "TRUE", "LLP_REPLACEMENT", "ENG1 HPT LLP"],
        ["ENG2_LLP_HPT", "Engine 2 HPT Disk (LLP)", "LLP", "FC", 18000, "PCT", 0.90, "FC", 1500, "TRUE", "LLP_REPLACEMENT", "ENG2 HPT LLP"],
        ["APU", "APU Shop Visit", "APU", "FH", 12000, "PCT", 0.80, "NONE", "", "TRUE", "HEAVY", "APU SV"],
        ["LG", "Landing Gear Overhaul", "LANDING_GEAR", "CALENDAR_MONTHS", 120, "PCT", 0.85, "MONTHS", 6, "FALSE", "HEAVY", "LG Overhaul"],
    ]
    reg_atr = [
        ["AF_C", "Airframe C-Check", "AIRFRAME", "CALENDAR_MONTHS", 24, "PCT", 0.85, "MONTHS", 3, "FALSE", "LIGHT", "C-Check"],
        ["ENG1_HSI", "Engine 1 Hot Section Inspection", "ENGINE", "FH", 6000, "PCT", 0.80, "FH", 800, "TRUE", "LIGHT", "ENG1 HSI"],
        ["ENG2_HSI", "Engine 2 Hot Section Inspection", "ENGINE", "FH", 6000, "PCT", 0.80, "FH", 800, "TRUE", "LIGHT", "ENG2 HSI"],
        ["ENG1_OVH", "Engine 1 Overhaul", "ENGINE", "FH", 12000, "PCT", 0.80, "FH", 1500, "TRUE", "HEAVY", "ENG1 OVH"],
        ["ENG2_OVH", "Engine 2 Overhaul", "ENGINE", "FH", 12000, "PCT", 0.80, "FH", 1500, "TRUE", "HEAVY", "ENG2 OVH"],
        ["PROP1", "Propeller 1 Overhaul", "LANDING_GEAR", "CALENDAR_MONTHS", 72, "PCT", 0.85, "MONTHS", 4, "FALSE", "LIGHT", "PROP1 OVH"],
        ["LG", "Landing Gear Overhaul", "LANDING_GEAR", "CALENDAR_MONTHS", 96, "PCT", 0.85, "MONTHS", 4, "FALSE", "HEAVY", "LG Overhaul"],
    ]
    for fam, rows in (("A320neo", reg_a320), ("ATR72", reg_atr)):
        ws = wb.create_sheet(f"REG_{fam}")
        _write_rows(ws, reg_headers, rows, example_rows=len(rows))
        _add_dropdown(ws, 3, "Category")
        _add_dropdown(ws, 4, "Driver")
        _add_dropdown(ws, 6, "SoftThresholdType")
        _add_dropdown(ws, 8, "RedeliveryMinLifeType")
        _add_dropdown(ws, 10, "Bool")
        _add_dropdown(ws, 11, "WorkscopeClass")

    # ---- LINKS_<family> : SV dependency graph (one per type) ----
    link_headers = ["Source_Component_ID", "Target_Component_ID", "Link_Type", "Notes"]
    links_a320 = [
        ["ENG1_PR", "ENG1_LLP_HPT", "SOFT_BUNDLE", "If HPT LLP near limit when ENG1 goes to shop, replace in same visit."],
        ["ENG2_PR", "ENG2_LLP_HPT", "SOFT_BUNDLE", "Same for engine 2."],
        ["AF_12Y", "AF_6Y", "HARD", "A 12Y check always encompasses the 6Y scope -> reset both."],
    ]
    links_atr = [
        ["ENG1_OVH", "ENG1_HSI", "HARD", "Full overhaul includes hot-section scope -> reset HSI too."],
        ["ENG2_OVH", "ENG2_HSI", "HARD", "Same for engine 2."],
    ]
    for fam, rows in (("A320neo", links_a320), ("ATR72", links_atr)):
        ws = wb.create_sheet(f"LINKS_{fam}")
        _write_rows(ws, link_headers, rows, example_rows=len(rows))
        _add_dropdown(ws, 3, "LinkType")
        ws.column_dimensions["D"].width = 60

    # ---- Fleet : one row per tail (the hub) ----
    fleet_headers = [
        "Tail", "Aircraft_Family", "MSN", "Operator", "Owner", "Status",
        "Delivery_Date", "Redelivery_Date", "Lease_Extension_Date",
        "Utilization_Profile_ID", "Seasonality_Profile_ID", "Financial_Scenario_ID",
    ]
    fleet_rows = [
        ["SX-NEA", "A320neo", "10231", "ExampleAir", "LessorCo", "ACTIVE", "2019-03-15", "2031-03-14", "", "UTIL_HIGH", "SEAS_MED", "FIN_LEASE_MR"],
        ["SX-NEB", "A320neo", "10488", "ExampleAir", "LessorCo", "ACTIVE", "2020-06-01", "2032-05-31", "", "UTIL_MED", "SEAS_MED", "FIN_LEASE_MR"],
        ["SX-ATA", "ATR72", "1455", "ExampleAir", "OwnedCo", "ACTIVE", "2017-09-10", "2027-09-09", "", "UTIL_ATR", "SEAS_ISL", "FIN_TM_OWNED"],
    ]
    ws = wb.create_sheet("Fleet")
    _write_rows(ws, fleet_headers, fleet_rows, example_rows=len(fleet_rows))
    _add_dropdown(ws, 6, "AircraftStatus")

    # ---- Initial_Status : one row per tail x component ----
    is_headers = ["Tail", "Component_ID", "TSN", "CSN", "SLSV_FH", "SLSV_FC", "SLSV_Months", "Last_SV_Date"]
    is_rows = [
        ["SX-NEA", "AF_6Y", 24500, 14200, "", "", 40, "2022-09-30"],
        ["SX-NEA", "ENG1_PR", 24500, 14200, 11800, 7100, "", "2023-01-31"],
        ["SX-NEA", "ENG1_LLP_HPT", 24500, 14200, "", 7100, "", "2023-01-31"],
        ["SX-ATA", "AF_C", 18900, 21300, "", "", 11, "2025-02-28"],
        ["SX-ATA", "ENG1_HSI", 18900, 21300, 3200, "", "", "2024-11-30"],
    ]
    ws = wb.create_sheet("Initial_Status")
    _write_rows(ws, is_headers, is_rows, example_rows=len(is_rows))

    # ---- Component_Limit_Overrides : per-tail hard-limit overrides (optional) ----
    clo_headers = ["Tail", "Component_ID", "Hard_Limit"]
    clo_rows = [["SX-NEA", "ENG1_LLP_HPT", 17200]]
    ws = wb.create_sheet("Component_Limit_Overrides")
    _write_rows(ws, clo_headers, clo_rows, example_rows=len(clo_rows))

    # ---- Utilization : reusable profiles ----
    util_headers = [
        "Utilization_Profile_ID", "Base_Monthly_FH", "Base_Monthly_FC",
        "Annual_Evolution_Pct", "Min_Monthly_FH", "Max_Monthly_FH",
        "Min_Monthly_FC", "Max_Monthly_FC", "AOG_Lead_In_Months",
    ]
    util_rows = [
        ["UTIL_HIGH", 300, 120, 0.02, 180, 360, 70, 150, 0],
        ["UTIL_MED", 230, 95, 0.015, 140, 290, 55, 120, 0],
        ["UTIL_ATR", 180, 200, 0.01, 110, 230, 120, 260, 1],
    ]
    ws = wb.create_sheet("Utilization")
    _write_rows(ws, util_headers, util_rows, example_rows=len(util_rows))

    # ---- Seasonality : 12 rows per profile ----
    seas_headers = ["Seasonality_Profile_ID", "Month", "FH_Factor", "FC_Factor"]
    # MED: summer-peaked (Greece mainland); ISL: strong summer island pattern
    med = [1.00, 0.95, 1.00, 1.05, 1.10, 1.15, 1.20, 1.20, 1.10, 1.00, 0.95, 0.90]
    isl = [0.70, 0.70, 0.85, 1.00, 1.20, 1.35, 1.45, 1.45, 1.20, 0.95, 0.75, 0.70]
    seas_rows = []
    for m in range(1, 13):
        seas_rows.append(["SEAS_MED", m, round(med[m - 1], 2), round(med[m - 1], 2)])
    for m in range(1, 13):
        seas_rows.append(["SEAS_ISL", m, round(isl[m - 1], 2), round(isl[m - 1], 2)])
    ws = wb.create_sheet("Seasonality")
    _write_rows(ws, seas_headers, seas_rows, example_rows=len(seas_rows))

    # ---- Financial_Scenarios : contract header ----
    fs_headers = ["Financial_Scenario_ID", "Name", "Currency", "Escalation_Base_Date", "Notes"]
    fs_rows = [
        ["FIN_LEASE_MR", "A320neo lease w/ maintenance reserves", "USD", "2019-01-01",
         "Engines on FH/FC-ratio MR, LLPs on EOL pro-rata, airframe on T&M."],
        ["FIN_TM_OWNED", "ATR72 owned — pure Time & Material", "USD", "2017-01-01",
         "All components costed as T&M shop visits (no reserves)."],
    ]
    ws = wb.create_sheet("Financial_Scenarios")
    _write_rows(ws, fs_headers, fs_rows, example_rows=len(fs_rows))

    # ---- Financial_Rules : Scenario -> scope -> method + params ----
    fr_headers = ["Scenario_ID", "Scope_Type", "Scope_Value", "Method_Type", "Param_Set_ID"]
    fr_rows = [
        ["FIN_LEASE_MR", "CATEGORY", "ENGINE", "MR_FH_FC_RATIO", "P_ENG_MR"],
        ["FIN_LEASE_MR", "CATEGORY", "LLP", "EOL_PRORATA", "P_LLP_EOL"],
        ["FIN_LEASE_MR", "CATEGORY", "AIRFRAME", "TM_WORKSCOPE", "P_AF_TM"],
        ["FIN_LEASE_MR", "COMPONENT", "APU", "PBH_FLAT", "P_APU_PBH"],
        ["FIN_TM_OWNED", "ALL", "", "TM_WORKSCOPE", "P_OWNED_TM"],
    ]
    ws = wb.create_sheet("Financial_Rules")
    _write_rows(ws, fr_headers, fr_rows, example_rows=len(fr_rows))
    _add_dropdown(ws, 2, "ScopeType")
    _add_dropdown(ws, 4, "MethodType")

    # ---- MR_Rates (params for MR_* methods) ----
    mr_headers = [
        "Param_Set_ID", "Component_Scope_Type", "Component_Scope_Value",
        "Rate_Basis", "Rate", "Escalation_Pct", "Reimbursable", "Reimbursement_Cap",
    ]
    mr_rows = [
        ["P_ENG_MR", "CATEGORY", "ENGINE", "PER_FH", 210.0, 0.03, "TRUE", 4500000],
        ["P_ENG_MR", "CATEGORY", "ENGINE", "PER_FC", 95.0, 0.03, "TRUE", 4500000],
    ]
    ws = wb.create_sheet("MR_Rates")
    _write_rows(ws, mr_headers, mr_rows, example_rows=len(mr_rows))
    _add_dropdown(ws, 2, "ScopeType")
    _add_dropdown(ws, 4, "RateBasis")
    _add_dropdown(ws, 7, "Bool")

    # ---- SV_Cost_Matrix (params for TM_WORKSCOPE / PBH) ----
    sc_headers = [
        "Param_Set_ID", "Workscope_Class", "Event_Type",
        "Fixed_Cost", "Cost_Per_FH", "LLP_Replacement_Cost",
    ]
    sc_rows = [
        ["P_AF_TM", "HEAVY", "HEAVY_SV", 850000, 0, 0],
        ["P_AF_TM", "LIGHT", "LIGHT_SV", 180000, 0, 0],
        ["P_OWNED_TM", "HEAVY", "HEAVY_SV", 1200000, 0, 0],
        ["P_OWNED_TM", "LIGHT", "LIGHT_SV", 240000, 0, 0],
        ["P_OWNED_TM", "LLP_REPLACEMENT", "LLP_REPLACEMENT", 0, 0, 2100000],
    ]
    ws = wb.create_sheet("SV_Cost_Matrix")
    _write_rows(ws, sc_headers, sc_rows, example_rows=len(sc_rows))
    _add_dropdown(ws, 2, "WorkscopeClass")
    _add_dropdown(ws, 3, "EventType")

    # ---- EOL_Terms (params for EOL_PRORATA) ----
    eol_headers = ["Param_Set_ID", "Component_Scope_Type", "Component_Scope_Value", "Comp_Basis", "Rate"]
    eol_rows = [
        ["P_LLP_EOL", "CATEGORY", "LLP", "PER_CONSUMED_FC", 140.0],
    ]
    ws = wb.create_sheet("EOL_Terms")
    _write_rows(ws, eol_headers, eol_rows, example_rows=len(eol_rows))
    _add_dropdown(ws, 2, "ScopeType")
    _add_dropdown(ws, 4, "CompBasis")

    # ---- Financial_Params : generic key-value knobs (the escape hatch) ----
    fp_headers = ["Param_Set_ID", "Param_Name", "Param_Value"]
    fp_rows = [
        ["P_APU_PBH", "pbh_rate_per_fh", 32.0],
        ["P_APU_PBH", "escalation_pct", 0.025],
        ["P_ENG_MR", "fh_fc_ratio_target", 2.3],
    ]
    ws = wb.create_sheet("Financial_Params")
    _write_rows(ws, fp_headers, fp_rows, example_rows=len(fp_rows))

    # ---- Save ----
    out = Path(__file__).resolve().parents[1] / "data" / "FLEET_ERP_DATABASE_TEMPLATE.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _build_instructions(ws):
    ws.column_dimensions["A"].width = 100
    ws.sheet_view.showGridLines = False

    lines = [
        ("Fleet ERP — Database Workbook", TITLE_FONT),
        ("This single file IS your database. Fill it in, upload it to the app, and re-upload an "
         "updated copy each cycle. The app never keeps a hidden copy — you always own the data.", NOTE_FONT),
        ("", None),
        ("How the workbook is organised (4 layers):", Font(bold=True, size=12, color="1F3864")),
        ("A. SETUP   — Config: one sheet, global run settings.", NOTE_FONT),
        ("B. TYPE LIBRARY — REG_<family> and LINKS_<family>: one pair PER aircraft type. "
         "Grows with number of TYPES, not aircraft.", NOTE_FONT),
        ("C. FLEET   — Fleet / Initial_Status / Component_Limit_Overrides: your aircraft live here "
         "as ROWS. Add 100 tails = paste 100 rows, no new sheets.", NOTE_FONT),
        ("D. SETTINGS / PROFILES — Utilization, Seasonality, Financial_* : defined ONCE and "
         "referenced by many tails via IDs.", NOTE_FONT),
        ("", None),
        ("The golden rule: DEFINE ONCE, REFERENCE BY ID.", Font(bold=True, size=12, color="C00000")),
        ("Each Fleet row points at everything else by ID: Aircraft_Family -> REG_/LINKS_ sheets; "
         "Utilization_Profile_ID -> Utilization; Seasonality_Profile_ID -> Seasonality; "
         "Financial_Scenario_ID -> Financial_Scenarios.", NOTE_FONT),
        ("", None),
        ("Adding things:", Font(bold=True, size=12, color="1F3864")),
        ("• New aircraft TYPE  -> duplicate a REG_ and LINKS_ sheet, rename the suffix, fill in its rows.", NOTE_FONT),
        ("• New aircraft       -> add a row to Fleet (+ its rows in Initial_Status).", NOTE_FONT),
        ("• New event to watch -> add a row to that type's REG_ sheet.", NOTE_FONT),
        ("• New dependency     -> add a row to that type's LINKS_ sheet (Source -> Target, HARD or SOFT_BUNDLE).", NOTE_FONT),
        ("• New contract       -> add a Financial_Scenarios row + Financial_Rules rows + parameters.", NOTE_FONT),
        ("", None),
        ("Financial modularity (Scenario -> Rules -> Method + Params):", Font(bold=True, size=12, color="1F3864")),
        ("A contract (Financial_Scenarios) is a set of Financial_Rules. Each rule routes a slice of "
         "components (ALL / a CATEGORY / one COMPONENT) to a Method_Type and a Param_Set_ID.", NOTE_FONT),
        ("One scenario can mix methods: see FIN_LEASE_MR — engines on MR_FH_FC_RATIO, LLPs on "
         "EOL_PRORATA, airframe on TM_WORKSCOPE, APU on PBH_FLAT.", NOTE_FONT),
        ("Method_Type values are calculation plugins in the app. Re-using an existing method for a new "
         "contract = pure data entry. A brand-new way of calculating = we add one plugin, then you can "
         "reference it from Excel forever after.", NOTE_FONT),
        ("", None),
        ("Conventions:", Font(bold=True, size=12, color="1F3864")),
        ("• Grey italic rows are EXAMPLES — overwrite or delete them.", EXAMPLE_FONT),
        ("• Dropdown cells only accept listed values (keeps data clean).", NOTE_FONT),
        ("• Dates: YYYY-MM-DD, month-end where possible. Leave cells blank rather than 0 when N/A.", NOTE_FONT),
        ("• Component_IDs must be UNIQUE within a family and must match exactly across REG_, LINKS_, "
         "Initial_Status and Financial scope values.", NOTE_FONT),
        ("• Do not rename sheets except the <family> suffix on REG_/LINKS_. Do not edit Schema_Version.", NOTE_FONT),
    ]
    r = 1
    for text, font in lines:
        c = ws.cell(row=r, column=1, value=text)
        if font:
            c.font = font
        c.alignment = WRAP_TOP
        if text == "":
            ws.row_dimensions[r].height = 6
        elif len(text) > 90:
            ws.row_dimensions[r].height = 30
        r += 1


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
